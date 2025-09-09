import os
import uuid
import io  # MISSING IMPORT - ADD THIS
from typing import List, Dict, Any, Optional
import fitz  # PyMuPDF
import pdfplumber
from docx import Document as DocxDocument
from pptx import Presentation
import openpyxl
from PIL import Image
import pytesseract
import mimetypes
from pathlib import Path
import tiktoken
from langchain.text_splitter import RecursiveCharacterTextSplitter

class DocumentProcessor:
    def __init__(self, upload_dir: str = "uploads"):
        self.upload_dir = Path(upload_dir)
        self.upload_dir.mkdir(exist_ok=True)
        
        # Initialize text splitter
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            length_function=len,
            separators=["\n\n", "\n", ". ", " ", ""]
        )
        
        # Initialize tokenizer for accurate token counting
        try:
            self.tokenizer = tiktoken.get_encoding("cl100k_base")
        except:
            self.tokenizer = None
    
    def count_tokens(self, text: str) -> int:
        """Count tokens in text"""
        if self.tokenizer:
            return len(self.tokenizer.encode(text))
        else:
            # Fallback approximation
            return len(text.split()) * 1.3
    
    async def save_uploaded_file(self, file_content: bytes, filename: str) -> tuple[str, str]:
        """Save uploaded file and return document_id and file_path"""
        document_id = str(uuid.uuid4())
        file_extension = Path(filename).suffix
        safe_filename = f"{document_id}{file_extension}"
        file_path = self.upload_dir / safe_filename
        
        with open(file_path, "wb") as f:
            f.write(file_content)
        
        return document_id, str(file_path)
    
    def extract_text_from_pdf(self, file_path: str) -> Dict[str, Any]:
        """Extract text from PDF using multiple methods with better error handling"""
        text_content = []
        metadata = {"pages": 0, "method": "mixed"}
        
        print(f"🔍 Extracting text from PDF: {file_path}")
        
        try:
            # Try PyMuPDF first (faster)
            doc = fitz.open(file_path)
            metadata["pages"] = len(doc)
            
            print(f"📄 PDF has {len(doc)} pages")
            
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text()
                
                print(f"📝 Page {page_num + 1}: extracted {len(text)} characters")
                
                # If no text found, try OCR on images
                if not text.strip():
                    print(f"🖼️ No text found on page {page_num + 1}, trying OCR...")
                    try:
                        pix = page.get_pixmap()
                        img_data = pix.tobytes("ppm")
                        text = pytesseract.image_to_string(Image.open(io.BytesIO(img_data)))
                        metadata["method"] = "ocr"
                        print(f"🔍 OCR extracted {len(text)} characters")
                    except Exception as ocr_error:
                        print(f"❌ OCR failed: {ocr_error}")
                        text = ""
                
                if text.strip():
                    text_content.append({
                        "page": page_num + 1,
                        "content": text.strip()
                    })
                else:
                    print(f"⚠️ No text found on page {page_num + 1}")
            
            doc.close()
            
        except Exception as e:
            print(f"❌ PyMuPDF failed: {e}, trying pdfplumber...")
            
            # Fallback to pdfplumber
            try:
                with pdfplumber.open(file_path) as pdf:
                    metadata["pages"] = len(pdf.pages)
                    print(f"📄 pdfplumber: PDF has {len(pdf.pages)} pages")
                    
                    for page_num, page in enumerate(pdf.pages):
                        text = page.extract_text()
                        print(f"📝 pdfplumber Page {page_num + 1}: extracted {len(text) if text else 0} characters")
                        
                        if text and text.strip():
                            text_content.append({
                                "page": page_num + 1,
                                "content": text.strip()
                            })
            except Exception as e:
                raise Exception(f"Failed to extract PDF text with both methods: {e}")
        
        # Check if we extracted any text
        total_text = sum(len(item["content"]) for item in text_content)
        print(f"✅ Total text extracted: {total_text} characters from {len(text_content)} pages")
        
        if total_text == 0:
            raise Exception("No text could be extracted from this PDF. It might be image-based, encrypted, or corrupted.")
        
        return {"content": text_content, "metadata": metadata}

    
    async def process_document(self, file_path: str, filename: str) -> Dict[str, Any]:
        """Main document processing method with better debugging"""
        file_extension = Path(filename).suffix.lower()
        
        print(f"🔧 Processing document: {filename} (type: {file_extension})")
        
        # Extract text based on file type
        if file_extension == '.pdf':
            extracted = self.extract_text_from_pdf(file_path)
        elif file_extension == '.docx':
            extracted = self.extract_text_from_docx(file_path)
        elif file_extension == '.pptx':
            extracted = self.extract_text_from_pptx(file_path)
        elif file_extension == '.xlsx':
            extracted = self.extract_text_from_xlsx(file_path)
        elif file_extension == '.txt':
            extracted = self.extract_text_from_txt(file_path)
        else:
            raise Exception(f"Unsupported file type: {file_extension}")
        
        # Combine all text content
        full_text = ""
        for content_block in extracted["content"]:
            full_text += content_block["content"] + "\n\n"
        
        print(f"📋 Combined text length: {len(full_text)} characters")
        
        if not full_text.strip():
            raise Exception("No text content found in document after extraction")
        
        # Create chunks
        chunks = self.text_splitter.split_text(full_text.strip())
        print(f"✂️ Created {len(chunks)} chunks")
        
        if len(chunks) == 0:
            raise Exception("Text splitter produced no chunks")
        
        # Add metadata to each chunk
        processed_chunks = []
        for i, chunk in enumerate(chunks):
            if not chunk.strip():  # Skip empty chunks
                continue
                
            chunk_data = {
                "chunk_index": i,
                "content": chunk.strip(),
                "token_count": self.count_tokens(chunk),
                "metadata": {
                    **extracted["metadata"],
                    "filename": filename,
                    "chunk_index": i,
                    "total_chunks": len(chunks)
                }
            }
            processed_chunks.append(chunk_data)
        
        print(f"✅ Processed {len(processed_chunks)} non-empty chunks")
        
        if len(processed_chunks) == 0:
            raise Exception("No non-empty chunks were created from the document")
        
        return {
            "chunks": processed_chunks,
            "total_chunks": len(processed_chunks),
            "metadata": extracted["metadata"]
        }

    def extract_text_from_docx(self, file_path: str) -> Dict[str, Any]:
        """Extract text from Word document"""
        try:
            doc = DocxDocument(file_path)
            text_content = []
            
            for para_num, para in enumerate(doc.paragraphs):
                if para.text.strip():
                    text_content.append({
                        "paragraph": para_num + 1,
                        "content": para.text.strip()
                    })
            
            metadata = {
                "paragraphs": len([p for p in doc.paragraphs if p.text.strip()]),
                "method": "docx"
            }
            
            return {"content": text_content, "metadata": metadata}
        
        except Exception as e:
            raise Exception(f"Failed to extract DOCX text: {e}")
    
    def extract_text_from_pptx(self, file_path: str) -> Dict[str, Any]:
        """Extract text from PowerPoint presentation"""
        try:
            prs = Presentation(file_path)
            text_content = []
            
            for slide_num, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                
                if slide_text:
                    text_content.append({
                        "slide": slide_num + 1,
                        "content": "\n".join(slide_text)
                    })
            
            metadata = {
                "slides": len(prs.slides),
                "method": "pptx"
            }
            
            return {"content": text_content, "metadata": metadata}
        
        except Exception as e:
            raise Exception(f"Failed to extract PPTX text: {e}")
    
    def extract_text_from_xlsx(self, file_path: str) -> Dict[str, Any]:
        """Extract text from Excel file"""
        try:
            wb = openpyxl.load_workbook(file_path)
            text_content = []
            
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_content = []
                
                for row in sheet.iter_rows(values_only=True):
                    row_data = [str(cell) for cell in row if cell is not None]
                    if row_data:
                        sheet_content.append(" | ".join(row_data))
                
                if sheet_content:
                    text_content.append({
                        "sheet": sheet_name,
                        "content": "\n".join(sheet_content)
                    })
            
            metadata = {
                "sheets": len(wb.sheetnames),
                "method": "xlsx"
            }
            
            return {"content": text_content, "metadata": metadata}
        
        except Exception as e:
            raise Exception(f"Failed to extract XLSX text: {e}")
    
    def extract_text_from_txt(self, file_path: str) -> Dict[str, Any]:
        """Extract text from plain text file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            text_content = [{"content": content}]
            metadata = {
                "characters": len(content),
                "method": "txt"
            }
            
            return {"content": text_content, "metadata": metadata}
        
        except Exception as e:
            raise Exception(f"Failed to extract TXT text: {e}")

    def get_supported_extensions(self) -> List[str]:
        """Get list of supported file extensions"""
        return ['.pdf', '.docx', '.pptx', '.xlsx', '.txt']
    
    def is_supported_file(self, filename: str) -> bool:
        """Check if file type is supported"""
        return Path(filename).suffix.lower() in self.get_supported_extensions()